import io
import re
from datetime import datetime
from html.parser import HTMLParser
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile

import pdfplumber
from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from pdfminer.pdfdocument import PDFPasswordIncorrect
from pdfminer.pdfparser import PDFSyntaxError
from pdfplumber.utils.exceptions import PdfminerException

MAX_RESUME_BYTES = 10 * 1024 * 1024


class ResumeHTMLParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.ignored_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in {"script", "style"}:
            self.ignored_depth += 1
        elif not self.ignored_depth and tag in {"br", "div", "li", "p", "section", "tr"}:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style"} and self.ignored_depth:
            self.ignored_depth -= 1
        elif not self.ignored_depth and tag in {"div", "li", "p", "section", "tr"}:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self.ignored_depth:
            self.parts.append(data)

app = FastAPI(title="Local Resume Parser")
app.add_middleware(
    CORSMiddleware,
    allow_origin_regex=r"^chrome-extension://[a-p]{32}$",
    allow_methods=["POST"],
    allow_headers=["content-type"],
)


def extract_text(data: bytes, filename: str) -> str:
    name = filename.lower()
    try:
        if name.endswith(".pdf"):
            with pdfplumber.open(io.BytesIO(data)) as pdf:
                return "\n".join(page.extract_text() or "" for page in pdf.pages)
        if name.endswith(".docx"):
            document = Document(io.BytesIO(data))
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            paragraphs.extend(
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            )
            return "\n".join(paragraphs)
        if name.endswith(".odt"):
            with ZipFile(io.BytesIO(data)) as archive:
                content = ElementTree.fromstring(archive.read("content.xml"))
            paragraphs = [
                "".join(element.itertext())
                for element in content.iter()
                if element.tag.rsplit("}", 1)[-1] in {"p", "h"}
            ]
            return "\n".join(paragraphs)
        if name.endswith((".html", ".htm")):
            parser = ResumeHTMLParser()
            parser.feed(data.decode("utf-8-sig"))
            return "".join(parser.parts)
        if name.endswith((".txt", ".md")):
            return data.decode("utf-8-sig")
    except (
        BadZipFile,
        ElementTree.ParseError,
        KeyError,
        OSError,
        PackageNotFoundError,
        PDFPasswordIncorrect,
        PDFSyntaxError,
        PdfminerException,
        ValueError,
    ) as exc:
        raise HTTPException(422, "The selected file could not be read.") from exc
    raise HTTPException(415, "Text extraction supports PDF, DOCX, ODT, HTML, TXT and MD resumes.")


def parse_profile(text: str) -> dict:
    lines = [line.strip(" \t•*-") for line in text.splitlines() if line.strip()]
    clean_lines = [re.sub(r"\s+", " ", line.replace("\ufffd", "-")).strip() for line in lines]
    profile = {
        "first_name": "",
        "last_name": "",
        "full_name": "",
        "email": "",
        "phone": "",
        "address": {"line1": "", "city": "", "state": "", "country": "", "zip": ""},
        "linkedin": "",
        "github": "",
        "website": "",
        "current_company": "",
        "current_title": "",
        "years_experience": "",
        "summary": "",
        "skills": [],
        "education": [],
        "experience": [],
        "latest_company": "",
        "latest_title": "",
        "certifications": [],
        "work_auth": {"authorized": "", "requires_sponsorship": ""},
        "willing_to_relocate": "",
        "salary_expectation": "",
        "notice_period": "",
        "eeo": {},
        "answers": {},
        "cv_text": text[:20000],
    }

    if clean_lines and re.fullmatch(r"[A-Za-z][A-Za-z'’-]*(?:\s+[A-Za-z][A-Za-z'’-]*){1,3}", clean_lines[0]):
        name = clean_lines[0]
        if not re.search(r"\b(engineer|developer|manager|designer|analyst|resume|curriculum vitae)\b", name, re.I):
            name_parts = name.split()
            profile["full_name"] = name
            profile["first_name"] = name_parts[0]
            profile["last_name"] = name_parts[-1]

    email = re.search(r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", text, re.I)
    if email:
        profile["email"] = email.group(0)

    phone = re.search(r"(?<!\w)(?:\+?\d[\d\s().-]{7,}\d)(?!\w)", text)
    if phone:
        profile["phone"] = re.sub(r"\s+", " ", phone.group(0)).strip()

    for key, pattern in (
        ("linkedin", r"(?:https?://)?(?:www\.)?linkedin\.com/[^\s|,)]+"),
        ("github", r"(?:https?://)?(?:www\.)?github\.com/[^\s|,)]+"),
    ):
        match = re.search(pattern, text, re.I)
        if match:
            url = match.group(0).rstrip(".,")
            profile[key] = url if url.lower().startswith("http") else f"https://{url}"
    website = re.search(r"https?://(?:www\.)?(?!linkedin\.com|github\.com)[^\s|,)]+", text, re.I)
    if website:
        profile["website"] = website.group(0).rstrip(".,")

    section_heading = re.compile(
        r"^(professional summary|summary|education|professional experience|work experience|employment history|"
        r"technical skills|skills|certifications|languages|projects)$",
        re.I,
    )
    sections: dict[str, list[str]] = {}
    current_section = ""
    for line in clean_lines:
        heading = line.strip(" :")
        if section_heading.fullmatch(heading):
            current_section = heading.lower()
            sections.setdefault(current_section, [])
        elif current_section:
            sections[current_section].append(line)

    skill_lines = sections.get("technical skills", sections.get("skills", []))
    if skill_lines:
        skills = []
        for skill_line in skill_lines:
            values = skill_line.split(":", 1)
            if len(values) == 2:
                skill_line = values[1]
            skills.extend(
                value.strip(" •-*")
                for value in re.split(r"[,;|•]", skill_line)
                if value.strip(" •-*")
            )
        profile["skills"] = list(dict.fromkeys(skills))[:100]

    profile["certifications"] = [
        line.strip(" •-*")
        for line in sections.get("certifications", [])
        if line.strip(" •-*")
    ]

    education_lines = sections.get("education", [])
    month_name = (
        r"(?:Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|Jun(?:e)?|"
        r"Jul(?:y)?|Aug(?:ust)?|Sep(?:t(?:ember)?)?|Oct(?:ober)?|"
        r"Nov(?:ember)?|Dec(?:ember)?)"
    )
    date_range = re.compile(
        rf"(?P<start>(?:{month_name}\s+)?\d{{4}})\s*[-–—]\s*"
        rf"(?P<end>(?:(?:{month_name}\s+)?\d{{4}}|present))",
        re.I,
    )

    def add_dates(record: dict, match: re.Match) -> None:
        start, end = match.group("start"), match.group("end")
        record["start_date"] = start
        record["end_date"] = end
        month_year = re.compile(
            r"(?P<month>[A-Za-z]+)\s+(?P<year>\d{4})|(?P<year_only>\d{4})",
            re.I,
        )

        def as_month(value: str) -> datetime | None:
            parsed = month_year.fullmatch(value.strip())
            if not parsed:
                return None
            if parsed.group("year_only"):
                return datetime(int(parsed.group("year_only")), 1, 1)
            month = parsed.group("month")[:3].title()
            try:
                return datetime.strptime(f"{month} {parsed.group('year')}", "%b %Y")
            except ValueError:
                return None

        start_month, end_month = as_month(start), as_month(end)
        if start_month and end_month and end.lower() != "present":
            months = max(0, (end_month.year - start_month.year) * 12 + end_month.month - start_month.month)
            record["duration_months"] = months
            years, remaining_months = divmod(months, 12)
            duration_parts = []
            if years:
                duration_parts.append(f"{years} year" + ("" if years == 1 else "s"))
            if remaining_months:
                duration_parts.append(f"{remaining_months} month" + ("" if remaining_months == 1 else "s"))
            record["duration"] = " ".join(duration_parts) or "0 months"
        else:
            record["duration"] = f"{start} - {end}"

    for index, line in enumerate(education_lines):
        match = date_range.search(line)
        if not match:
            continue
        degree = line[:match.start()].strip(" |,-")
        field_match = re.search(r"\bin\s+(.+?)(?:\s*\([^)]*\))?$", degree, re.I)
        record = {
            "degree": degree,
            "field": field_match.group(1).strip() if field_match else "",
            "school": "",
        }
        add_dates(record, match)
        if index + 1 < len(education_lines):
            record["school"] = education_lines[index + 1].strip()
        if record["degree"] or record["school"]:
            profile["education"].append(record)

    experience_lines = sections.get("professional experience", sections.get("work experience", sections.get("employment history", [])))
    for index, line in enumerate(experience_lines):
        match = date_range.search(line)
        if not match:
            continue
        company_line = line[:match.start()].strip(" |,-")
        title = experience_lines[index - 1] if index else ""
        if title.startswith(("•", "-", "*")):
            title = ""
        company = company_line
        location = ""
        if "," in company_line:
            company, location = (part.strip() for part in company_line.rsplit(",", 1))
        record = {
            "title": title,
            "company": company,
            "location": location,
            "description": "",
        }
        add_dates(record, match)
        next_entry = next(
            (position for position in range(index + 1, len(experience_lines)) if date_range.search(experience_lines[position])),
            len(experience_lines),
        )
        details_end = next_entry - 1 if next_entry < len(experience_lines) else next_entry
        details = experience_lines[index + 1:details_end]
        record["description"] = " ".join(details).strip()
        if not profile["latest_company"]:
            profile["latest_company"] = company
            profile["latest_title"] = title
        if match.group("end").lower() == "present":
            profile["current_company"] = company
            profile["current_title"] = title
        profile["experience"].append(record)

    summary_lines = sections.get("professional summary", sections.get("summary", []))
    if summary_lines:
        profile["summary"] = " ".join(summary_lines).strip()
    if not profile["years_experience"]:
        years = re.search(r"\b(\d+(?:\.\d+)?)\s*\+?\s+years?\s+of\s+(?:hands-on\s+)?experience\b", text, re.I)
        if years:
            profile["years_experience"] = years.group(1)

    return profile


@app.get("/health")
def health():
    return {"ok": True, "mode": "local-only"}


@app.post("/parse-cv")
async def parse_cv(file: UploadFile = File(...)):
    filename = file.filename or ""
    data = await file.read(MAX_RESUME_BYTES + 1)
    if len(data) > MAX_RESUME_BYTES:
        raise HTTPException(413, "Resume files must be 10 MB or smaller.")

    text = extract_text(data, filename).strip()
    if len(text) < 20:
        raise HTTPException(
            422,
            "No readable text found. Scanned/image-only PDFs are not supported.",
        )
    return parse_profile(text)
